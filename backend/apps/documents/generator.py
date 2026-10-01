from io import BytesIO

from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from apps.common.choices import ControlledDocumentKind


POE_SECTIONS = [
    "Objetivo", "Alcance", "Responsabilidades", "Definiciones",
    "Recursos", "Desarrollo del procedimiento", "Registros", "Referencias",
]
MAX_ROWS = 500
MAX_COLUMNS = 20


def initial_content(kind, *, project, requirement=None):
    if kind == ControlledDocumentKind.PROCEDURE:
        values = {
            "Alcance": project.scope,
            "Referencias": project.standard.name + (
                f" - {requirement.clause}: {requirement.title}" if requirement else ""
            ),
        }
        return {"sections": [{"title": title, "text": values.get(title, "")} for title in POE_SECTIONS]}
    columns = (
        ["Requisito", "Control", "Documento asociado", "Responsable", "Fecha prevista", "Estado"]
        if kind == ControlledDocumentKind.MATRIX
        else ["Fecha", "Actividad", "Responsable", "Resultado", "Observaciones"]
    )
    return {"columns": columns, "rows": []}


def validate_content(kind, content):
    if not isinstance(content, dict):
        raise ValueError("El contenido debe ser un objeto.")
    if kind == ControlledDocumentKind.PROCEDURE:
        sections = content.get("sections")
        if not isinstance(sections, list) or not 1 <= len(sections) <= 30:
            raise ValueError("El POE necesita entre 1 y 30 secciones.")
        normalized = []
        for section in sections:
            if not isinstance(section, dict):
                raise ValueError("Cada seccion debe tener titulo y texto.")
            title, text = section.get("title"), section.get("text", "")
            if not isinstance(title, str) or not title.strip() or len(title) > 120:
                raise ValueError("Cada seccion necesita un titulo de hasta 120 caracteres.")
            if not isinstance(text, str) or len(text) > 20000:
                raise ValueError("El texto de cada seccion debe tener hasta 20000 caracteres.")
            normalized.append({"title": title.strip(), "text": text.strip()})
        return {"sections": normalized}
    columns, rows = content.get("columns"), content.get("rows", [])
    if not isinstance(columns, list) or not 1 <= len(columns) <= MAX_COLUMNS:
        raise ValueError(f"La hoja necesita entre 1 y {MAX_COLUMNS} columnas.")
    if any(not isinstance(c, str) or not c.strip() or len(c) > 100 for c in columns):
        raise ValueError("Las columnas necesitan un nombre de hasta 100 caracteres.")
    columns = [c.strip() for c in columns]
    if len({c.casefold() for c in columns}) != len(columns):
        raise ValueError("Los nombres de las columnas no pueden repetirse.")
    if not isinstance(rows, list) or len(rows) > MAX_ROWS:
        raise ValueError(f"La hoja admite hasta {MAX_ROWS} filas.")
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns):
            raise ValueError("Cada fila debe tener una celda por columna.")
        if any(not isinstance(cell, str) or len(cell) > 2000 for cell in row):
            raise ValueError("Cada celda debe ser texto de hasta 2000 caracteres.")
    return {"columns": columns, "rows": [[c.strip() for c in row] for row in rows]}


def render_document(controlled, number, content, *, created_by, approved_by=None, effective_date=None):
    approved = approved_by is not None
    metadata = [
        ("Empresa", controlled.project.company.name),
        ("Proyecto", controlled.project.name),
        ("Codigo", controlled.code),
        ("Version", f"{number:02d}"),
        ("Proceso", controlled.process_area or "Sin asignar"),
        ("Norma", controlled.project.standard.name),
        ("Requisito", controlled.requirement.clause if controlled.requirement_id else "Sin asociar"),
        ("Estado al emitir", "VIGENTE" if approved else "BORRADOR"),
        ("Elaborado por", created_by.get_username() if created_by else "Sin asignar"),
    ]
    if approved:
        metadata += [
            ("Aprobado por", approved_by.get_username()),
            ("Fecha de vigencia", effective_date.isoformat()),
        ]
    suffix = ".docx" if controlled.kind == ControlledDocumentKind.PROCEDURE else ".xlsx"
    payload = (
        _render_word(controlled.title, content, metadata)
        if suffix == ".docx" else _render_workbook(controlled.title, content, metadata)
    )
    return f"{controlled.code}_v{number:02d}{suffix}", payload


def _render_word(title, content, metadata):
    doc = WordDocument()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name in ["Heading 1", "Heading 2"]:
        doc.styles[name].font.color.rgb = RGBColor.from_string("176B61")
    doc.add_heading(title, 0)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    doc.add_paragraph()
    for index, item in enumerate(content["sections"], 1):
        doc.add_heading(f"{index}. {item['title']}", 1)
        for line in (item["text"] or "[Pendiente de completar]").splitlines():
            doc.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.add_run(f"{dict(metadata)['Codigo']} | v{dict(metadata)['Version']} | Pagina ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = title
    doc.core_properties.author = dict(metadata)["Elaborado por"]
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _literal_cell(sheet, row, column, value):
    cell = sheet.cell(row=row, column=column, value=value)
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _render_workbook(title, content, metadata):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contenido"
    column_count = len(content["columns"])
    last_column = get_column_letter(column_count)
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=column_count)
    cell = _literal_cell(sheet, 1, 1, title)
    cell.font = Font(name="Calibri", size=16, bold=True, color="176B61")
    sheet.row_dimensions[1].height = 26
    for i, heading in enumerate(content["columns"], 1):
        _literal_cell(sheet, 4, i, heading)
        sheet.column_dimensions[get_column_letter(i)].width = 25
    rows = content["rows"] or [[""] * column_count for _ in range(10)]
    for i, row in enumerate(rows, 5):
        sheet.row_dimensions[i].height = 32
        for j, value in enumerate(row, 1):
            _literal_cell(sheet, i, j, value)
    table = Table(displayName="TablaDocumento", ref=f"A4:{last_column}{4 + len(rows)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:4"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{4 + len(rows)}"
    control = workbook.create_sheet("Control documental")
    control.column_dimensions["A"].width = 25
    control.column_dimensions["B"].width = 65
    for i, (label, value) in enumerate(metadata, 1):
        _literal_cell(control, i, 1, label).font = Font(bold=True)
        _literal_cell(control, i, 2, value)
        control.row_dimensions[i].height = 28
    control.sheet_view.showGridLines = False
    workbook.properties.title = title
    workbook.properties.creator = dict(metadata)["Elaborado por"]
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
