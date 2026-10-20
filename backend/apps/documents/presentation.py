"""Deterministic document presentation. No AI output or compliance state changes."""
from io import BytesIO
from pathlib import Path
from math import ceil
from xml.sax.saxutils import escape

from django.core.files.storage import default_storage
from docx.shared import Cm, Pt
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, A3, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle


def logo_bytes(presentation):
    name = presentation.get("logo_name")
    if name and default_storage.exists(name):
        with default_storage.open(name, "rb") as source:
            return source.read()
    return None


def add_word_header(doc, presentation, metadata):
    section = doc.sections[0]
    section.top_margin = Cm(3.5)
    section.header_distance = Cm(.6)
    width = section.page_width - section.left_margin - section.right_margin
    table = section.header.add_table(rows=1, cols=2, width=width)
    table.columns[0].width, table.columns[1].width = Cm(4), width - Cm(4)
    logo = logo_bytes(presentation)
    if logo:
        from PIL import Image
        with Image.open(BytesIO(logo)) as img:
            w, h = img.size
        scale = min(3.2 / w, 1.4 / h)
        table.cell(0, 0).paragraphs[0].add_run().add_picture(BytesIO(logo), width=Cm(w * scale))
    header = presentation.get("header", {})
    table.cell(0, 1).text = "\n".join(filter(None, [
        header.get("title") or presentation.get("company_name", ""), header.get("subtitle", ""),
        f"{dict(metadata).get('Codigo', '')} · Versión {dict(metadata).get('Version', '')}",
    ]))
    for p in table.cell(0, 1).paragraphs:
        for run in p.runs:
            run.font.size = Pt(9)
    if header.get("footer"):
        section.footer.paragraphs[0].add_run(header["footer"] + "\n")


def literal_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, value if value is not None else "")
    if isinstance(value, str):
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def add_excel_header(sheet, title, metadata, presentation, columns, rows):
    """Common header for generated matrices and all report worksheets."""
    last = max(2, len(columns))
    header = presentation.get("header", {})
    for row, value in [
        (1, header.get("title") or presentation.get("company_name", "")),
        (2, title),
        (3, " | ".join(f"{k}: {v}" for k, v in metadata.items())),
        (4, header.get("subtitle", "")),
    ]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
        literal_cell(sheet, row, 1, value).font = Font(name="Calibri", size=14 if row < 3 else 10, bold=row < 3, color="173D65")
    sheet.row_dimensions[1].height = max(30, 18 * ceil(len(str(sheet["A1"].value or "")) / (last * 22)))
    sheet.row_dimensions[2].height = 28
    sheet.row_dimensions[3].height = max(44, 15 * ceil(len(str(sheet["A3"].value or "")) / (last * 22)))
    logo = logo_bytes(presentation)
    if logo:
        image = ExcelImage(BytesIO(logo))
        scale = min(100 / image.width, 45 / image.height)
        image.width, image.height = image.width * scale, image.height * scale
        sheet.add_image(image, f"{get_column_letter(last + 1)}1")
    for j, heading in enumerate(columns, 1):
        cell = literal_cell(sheet, 6, j, heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173D65")
        sheet.column_dimensions[get_column_letter(j)].width = 27
    sheet.row_dimensions[6].height = 32
    for i, row in enumerate(rows, 7):
        lines = max((sum(max(1, ceil(len(part) / 24)) for part in str(value).split("\n")) for value in row), default=1)
        sheet.row_dimensions[i].height = min(409, max(48, lines * 15))
        for j, value in enumerate(row, 1):
            cell = literal_cell(sheet, i, j, value)
            if i % 2:
                cell.fill = PatternFill("solid", fgColor="F0F5FA")
    end = max(6, 6 + len(rows))
    sheet.auto_filter.ref = f"A6:{get_column_letter(len(columns))}{end}"
    sheet.freeze_panes = "C7" if len(columns) > 2 else "A7"
    sheet.print_title_rows = "1:6"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if len(columns) > 10 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_area = f"A1:{get_column_letter(last + (1 if logo else 0))}{end}"
    sheet.oddHeader.center.text = '&B' + title.replace('&', '&&')[:150]
    sheet.oddFooter.left.text = header.get("footer", "").replace('&', '&&')[:180]
    sheet.oddFooter.right.text = "Página &P de &N"


def render_pdf(title, content, metadata, presentation):
    # Vera is shipped with ReportLab; this also works inside the Linux worker.
    import reportlab
    font_dir = Path(reportlab.__file__).parent / "fonts"
    for name, file in [("Norman", "Vera.ttf"), ("NormanBold", "VeraBd.ttf")]:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_dir / file)))
    matrix = "columns" in content
    page = landscape(A3 if len(content.get("columns", [])) > 10 else A4) if matrix else A4
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=page, topMargin=100, bottomMargin=48,
                            leftMargin=38, rightMargin=38, title=title,
                            author=dict(metadata).get("Elaborado por", "NORMAN"))
    body = ParagraphStyle("body", fontName="Norman", fontSize=9, leading=14, spaceAfter=8, alignment=TA_LEFT)
    heading = ParagraphStyle("heading", parent=body, fontName="NormanBold", fontSize=13, leading=17, spaceBefore=12)
    def p(value, style=body):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)
    story = [p(title, heading), p(" | ".join(f"{k}: {v}" for k, v in metadata)), Spacer(1, 12)]
    for index, section in enumerate(content.get("sections", []), 1):
        story.append(p(f"{index}. {section['title']}", heading))
        # Paragraphs keep arbitrarily long sections splittable across pages.
        for line in (section["text"] or "[Pendiente de completar]").splitlines():
            story.append(p(line or " "))
    if matrix:
        data = [[p(c) for c in content["columns"]]] + [[p(c) for c in row] for row in content["rows"]]
        table = LongTable(data, colWidths=[doc.width / len(content["columns"])] * len(content["columns"]), repeatRows=1, splitInRow=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DFEAF4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#B8C7D5")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        story.append(table)
    logo = logo_bytes(presentation)
    header = presentation.get("header", {})
    header_style = ParagraphStyle("header", parent=body, fontName="NormanBold", fontSize=10, leading=13)
    header_text = "\n".join(filter(None, [header.get("title") or presentation.get("company_name", ""), header.get("subtitle", ""),
        f"{dict(metadata).get('Codigo', '')} · Versión {dict(metadata).get('Version', '')} · Estado al emitir: {dict(metadata).get('Estado al emitir', '')}"]))
    header_left = 145 if logo else 38
    _, header_height = p(header_text, header_style).wrap(page[0] - header_left - 38, 200)
    doc.topMargin = max(100, 30 + header_height + 24)
    def decorate(canvas, document):
        width, height = page
        canvas.saveState()
        if logo:
            canvas.drawImage(ImageReader(BytesIO(logo)), 38, height - 72, width=90, height=42, preserveAspectRatio=True, mask="auto")
        left = 145 if logo else 38
        hp = p(header_text, header_style)
        _, h = hp.wrap(width - left - 38, 70)
        hp.drawOn(canvas, left, height - 30 - h)
        canvas.setStrokeColor(colors.HexColor("#173D65"))
        canvas.line(38, height - document.topMargin + 14, width - 38, height - document.topMargin + 14)
        canvas.setFont("Norman", 7)
        fp = p(header.get("footer", ""), ParagraphStyle("footer", parent=body, fontSize=7, leading=9))
        _, fh = fp.wrap(width - 160, 30)
        fp.drawOn(canvas, 38, max(14, 30 - fh))
        canvas.drawRightString(width - 38, 22, f"Página {document.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return output.getvalue()
