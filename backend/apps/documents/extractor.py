from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


class UnsupportedDocumentTypeError(ValueError):
    pass


def extract_document(file_path: Path) -> tuple[str, dict]:
    suffix = file_path.suffix.lower()
    if suffix == ".txt":
        return _extract_txt(file_path)
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix == ".xlsx":
        return _extract_xlsx(file_path)
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    raise UnsupportedDocumentTypeError(f"Unsupported file extension: {suffix}")


def _extract_txt(file_path: Path) -> tuple[str, dict]:
    encodings = ["utf-8-sig", "utf-8", "latin-1"]
    last_error = None
    for encoding in encodings:
        try:
            text = file_path.read_text(encoding=encoding)
            metadata = {
                "source_type": "txt",
                "encoding": encoding,
                "character_count": len(text),
                "line_count": len(text.splitlines()),
            }
            return text.strip(), metadata
        except UnicodeDecodeError as exc:
            last_error = exc
    raise last_error or ValueError("Unable to decode text file")


def _extract_docx(file_path: Path) -> tuple[str, dict]:
    document = DocxDocument(file_path)
    paragraph_lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_lines = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))

    content = paragraph_lines + table_lines
    text = "\n".join(content).strip()
    metadata = {
        "source_type": "docx",
        "paragraph_count": len(paragraph_lines),
        "table_count": len(document.tables),
        "table_row_count": len(table_lines),
        "character_count": len(text),
    }
    return text, metadata


def _extract_xlsx(file_path: Path) -> tuple[str, dict]:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    sheet_names = workbook.sheetnames
    lines = []
    non_empty_rows = 0
    try:
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            lines.append(f"# Sheet: {sheet_name}")
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value not in (None, "")]
                if values:
                    non_empty_rows += 1
                    lines.append("\t".join(values))
    finally:
        workbook.close()

    text = "\n".join(lines).strip()
    metadata = {
        "source_type": "xlsx",
        "sheet_count": len(sheet_names),
        "sheet_names": sheet_names,
        "non_empty_rows": non_empty_rows,
        "character_count": len(text),
    }
    return text, metadata


def _extract_pdf(file_path: Path) -> tuple[str, dict]:
    reader = PdfReader(str(file_path))
    page_texts = []
    pages_with_text = 0

    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        page_texts.append(text)
        if text:
            pages_with_text += 1

    joined_text = "\n\n".join(filter(None, page_texts)).strip()
    metadata = {
        "source_type": "pdf",
        "page_count": len(reader.pages),
        "pages_with_text": pages_with_text,
        "character_count": len(joined_text),
    }
    return joined_text, metadata
