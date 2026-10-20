from io import BytesIO
from pathlib import Path
import hashlib
import shutil
from uuid import uuid4

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from docx import Document as WordDocument
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from rest_framework.test import APITestCase

from apps.companies.models import Company, CompanyMembership, CompanyRole
from apps.documents.control import create_controlled_document, create_new_revision, transition_document
from apps.implementation.models import Project
from apps.reviews.report_generator import generate_project_progress_report
from apps.standards.models import Standard, StandardRequirement


class ExportTests(APITestCase):
    def setUp(self):
        root = Path(settings.BASE_DIR) / "test-media"
        root.mkdir(exist_ok=True)
        directory = root / f"exports-{uuid4().hex}"
        directory.mkdir()
        def cleanup():
            if not directory.resolve().is_relative_to(root.resolve()):
                raise RuntimeError("Directorio de pruebas fuera del proyecto")
            shutil.rmtree(directory)
        self.addCleanup(cleanup)
        override = override_settings(MEDIA_ROOT=str(directory))
        override.enable()
        self.addCleanup(override.disable)
        self.owner = get_user_model().objects.create_user(username="export-owner")
        self.reader = get_user_model().objects.create_user(username="export-reader")
        self.outsider = get_user_model().objects.create_user(username="export-outsider")
        self.company = Company.objects.create(name="Empresa de prueba", ruc="EXPORT-QA", created_by=self.owner)
        CompanyMembership.objects.create(company=self.company, user=self.reader, role=CompanyRole.READER)
        self.standard = Standard.objects.create(code="EXPORT-QA", name="Criterios ficticios")
        self.requirement = StandardRequirement.objects.create(standard=self.standard, clause="1", title="Control", requirement_text="Criterio ficticio", requires_real_evidence=True)
        self.project = Project.objects.create(company=self.company, standard=self.standard, name="Proyecto de prueba")
        self.client.force_authenticate(self.owner)

    def logo(self, color="blue"):
        buffer = BytesIO()
        Image.new("RGB", (240, 80), color).save(buffer, "PNG")
        return SimpleUploadedFile("logo.png", buffer.getvalue(), "image/png")

    def branding(self, title="Encabezado histórico", color="blue"):
        import json
        response = self.client.patch(f"/api/companies/{self.company.pk}/", {
            "logo": self.logo(color), "document_header": json.dumps({"title": title, "subtitle": "Calidad y gestión", "footer": "Copia de consulta"}),
        }, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)
        self.company.refresh_from_db()

    def controlled(self, kind="POE", content=None):
        return create_controlled_document(user=self.owner, project=self.project, code=f"QA-{kind}", title="Procedimiento de recepción", kind=kind,
            requirement=self.requirement, content=content or {"sections": [{"title": "Objetivo", "text": "Conservar registros de recepción."}]})

    def download(self, controlled, format, version=1):
        response = self.client.get(f"/api/controlled-documents/{controlled.pk}/download/?version={version}&file_format={format}")
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        payload = b"".join(response.streaming_content)
        self.assertEqual(response["Cache-Control"], "private, no-store")
        return payload

    def test_logo_and_header_validate_actual_image_and_multipart(self):
        self.branding()
        self.assertEqual(self.company.document_header["title"], "Encabezado histórico")
        response = self.client.get(f"/api/companies/{self.company.pk}/download/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Image.open(BytesIO(b"".join(response.streaming_content))).format, "PNG")
        for upload in [SimpleUploadedFile("fake.png", b"not an image", "image/png"), SimpleUploadedFile("large.png", b"x" * (2 * 1024 * 1024 + 1), "image/png")]:
            self.assertEqual(self.client.patch(f"/api/companies/{self.company.pk}/", {"logo": upload}, format="multipart").status_code, 400)
        for header in [{"unknown": "x"}, {"title": "x" * 121}, {"footer": 42}, {"title": "Inválido\u0001"}]:
            self.assertEqual(self.client.patch(f"/api/companies/{self.company.pk}/", {"document_header": header}, format="json").status_code, 400)

    def test_historical_header_and_logo_survive_company_changes_and_approval(self):
        self.branding()
        controlled = self.controlled()
        revision = controlled.revisions.get(number=1)
        logo_name = revision.presentation_snapshot["logo_name"]
        with revision.document.file.open("rb") as source:
            original = source.read()
        self.branding("Encabezado nuevo", "red")
        transition_document(controlled, user=self.owner, expected_version=1, action="submit")
        transition_document(controlled, user=self.owner, expected_version=1, action="approve", confirmed=True, notes="Revisión humana de prueba")
        revision.refresh_from_db()
        self.assertEqual(revision.presentation_snapshot["logo_name"], logo_name)
        self.assertTrue(revision.document.file.storage.exists(logo_name))
        with revision.document.file.open("rb") as source:
            self.assertEqual(source.read(), original)
        word = WordDocument(BytesIO(self.download(controlled, "docx")))
        header = " ".join(c.text for t in word.sections[0].header.tables for row in t.rows for c in row.cells)
        self.assertIn("Encabezado histórico", header)
        self.assertNotIn("Encabezado nuevo", header)
        pdf = PdfReader(BytesIO(self.download(controlled, "pdf")))
        self.assertIn("Encabezado histórico", pdf.pages[0].extract_text())
        next_revision = create_new_revision(controlled, user=self.owner, expected_version=1, change_summary="Nueva versión")
        self.assertEqual(next_revision.presentation_snapshot["header"]["title"], "Encabezado nuevo")

    def test_docx_pdf_and_xlsx_keep_unicode_text_and_literal_formulas(self):
        self.branding()
        content = {"columns": ["Control", "Responsable"], "rows": [["=HYPERLINK(\"bad\")", "María & José <QA>"], ["Recepción", "Gestión"]]}
        controlled = self.controlled("MATRIZ", content)
        workbook = load_workbook(BytesIO(self.download(controlled, "xlsx")))
        self.assertEqual(workbook["Contenido"]["A7"].data_type, "s")
        self.assertEqual(workbook["Contenido"].freeze_panes, "A7")
        self.assertEqual(workbook["Contenido"].print_title_rows, "$1:$6")
        word = WordDocument(BytesIO(self.download(controlled, "docx")))
        self.assertIn("María & José <QA>", " ".join(c.text for t in word.tables for row in t.rows for c in row.cells))
        pdf = PdfReader(BytesIO(self.download(controlled, "pdf")))
        self.assertIn("María & José <QA>", "".join(p.extract_text() for p in pdf.pages))

    def test_exports_and_logo_obey_company_roles_and_do_not_change_states(self):
        self.branding()
        controlled = self.controlled()
        before = list(self.project.checklist_items.values_list("status", flat=True))
        for user, expected in [(self.reader, 200), (self.outsider, 404)]:
            self.client.force_authenticate(user)
            for url in [f"/api/projects/{self.project.pk}/export_report/", f"/api/companies/{self.company.pk}/download/", f"/api/controlled-documents/{controlled.pk}/download/?version=1&file_format=pdf"]:
                response = self.client.get(url)
                self.assertEqual(response.status_code, expected)
                if response.streaming:
                    b"".join(response.streaming_content)
            self.assertIn(self.client.patch(f"/api/companies/{self.company.pk}/", {"document_header": {"title": "Forbidden"}}, format="json").status_code, [403, 404])
        self.assertEqual(list(self.project.checklist_items.values_list("status", flat=True)), before)

    def test_excel_report_includes_every_requirement_and_all_records(self):
        self.branding()
        for i in range(2, 27):
            StandardRequirement.objects.create(standard=self.standard, clause=str(i), title=f"Criterio {i}", requirement_text="Ficticio", sequence=i)
        from apps.implementation.services import generate_checklist_for_project
        generate_checklist_for_project(self.project)
        self.project.name = "=No ejecutar"
        self.project.save()
        response = self.client.get(f"/api/projects/{self.project.pk}/export_report/")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(workbook.sheetnames, ["Resumen", "Matriz de requisitos", "Acciones", "Evidencias", "Documentos", "Revisiones", "Evaluaciones"])
        matrix = workbook["Matriz de requisitos"]
        self.assertEqual(matrix.max_row, 32)
        self.assertEqual(matrix.freeze_panes, "C7")
        self.assertIn("=No ejecutar", matrix["A3"].value)
        self.assertEqual(matrix["A3"].data_type, "s")
        self.assertEqual(matrix.auto_filter.ref, "A6:M32")

    def test_uploaded_original_is_never_silently_converted(self):
        controlled = self.controlled()
        original = self.download(controlled, "docx")
        revision = create_new_revision(controlled, user=self.owner, expected_version=1, change_summary="Archivo externo", upload=SimpleUploadedFile("external.docx", original))
        response = self.client.get(f"/api/controlled-documents/{controlled.pk}/download/?version={revision.number}&file_format=pdf")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(hashlib.sha256(self.download(controlled, "docx", version=revision.number)).hexdigest(), revision.content_hash)

    def test_long_procedure_paginates_with_repeated_header(self):
        self.branding()
        controlled = self.controlled(content={"sections": [{"title": "Desarrollo", "text": "Recepción y trazabilidad. " * 700}]})
        pdf = PdfReader(BytesIO(self.download(controlled, "pdf")))
        self.assertGreater(len(pdf.pages), 1)
        self.assertTrue(all("Encabezado histórico" in page.extract_text() for page in pdf.pages))

    def test_documentary_validation_never_increases_real_implementation_percentage(self):
        item = self.project.checklist_items.get(requirement=self.requirement)
        item.status, item.progress_percentage = "VALIDADO_DOCUMENTALMENTE", 55
        item.save()
        summary = generate_project_progress_report(self.project)["summary"]
        self.assertEqual(summary["documentary_progress_percentage"], 100)
        self.assertEqual(summary["implementation_progress_percentage"], 0)
        item.status = "IMPLEMENTADO"
        item.save()
        self.assertEqual(generate_project_progress_report(self.project)["summary"]["implementation_progress_percentage"], 100)
